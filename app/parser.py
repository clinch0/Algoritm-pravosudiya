"""Извлечение текста из .pdf, .docx и .txt файлов."""
from __future__ import annotations
from pathlib import Path


def extract_text(path: str) -> str:
    """Вернуть текст документа. Поддерживаются .pdf, .docx, .txt."""
    p = Path(path)
    if not p.exists():
        raise ValueError(f"Файл не найден: {path}")
    ext = p.suffix.lower()
    from . import config
    with p.open("rb") as stream:
        content = stream.read(config.MAX_UPLOAD_MB * 1024 * 1024 + 1)
    return extract_bytes(content, ext)


def _from_pdf(p: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(p))
    parts = []
    for page in reader.pages:
        t = page.extract_text() or ""
        parts.append(t)
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("Не удалось извлечь текст из PDF (возможно, документ является скан-образом).")
    return text


def _from_docx(p: Path) -> str:
    from docx import Document
    doc = Document(str(p))
    parts = [para.text for para in doc.paragraphs if para.text.strip()]
    # таблицы тоже читаем
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("Не удалось извлечь текст из DOCX: документ пуст.")
    return text


def _from_txt(p: Path) -> str:
    for enc in ("utf-8", "cp1251"):
        try:
            text = p.read_text(encoding=enc)
            if text.strip():
                return text.strip()
        except UnicodeDecodeError:
            continue
    raise ValueError("Не удалось прочитать текстовый файл (кодировка не распознана).")


def extract_bytes(content: bytes, extension: str) -> str:
    """Parse an in-memory buffer; framework upload spooling is a separate concern."""
    from io import BytesIO
    from zipfile import ZipFile, BadZipFile
    from . import config
    if len(content) > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise ValueError("Превышен максимальный размер файла")
    extension = extension.lower()
    if extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ValueError("PDF защищён паролем; предоставьте доступную текстовую копию")
        if len(reader.pages) > 200:
            raise ValueError("Максимум 200 страниц PDF в демо")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif extension == ".docx":
        from docx import Document
        try:
            with ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError("Слишком большой распакованный DOCX")
        except BadZipFile:
            raise ValueError("Повреждённый DOCX") from None
        doc = Document(BytesIO(content))
        # Preserve paragraph/table order rather than appending all tables at the end.
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        parts = []
        for item in doc.iter_inner_content():
            if isinstance(item, Paragraph):
                parts.append(item.text)
            elif isinstance(item, Table):
                parts.extend(" | ".join(c.text for c in row.cells) for row in item.rows)
        text = "\n".join(parts)
    elif extension in (".txt", ".md"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1251")
    else:
        raise ValueError("Неподдерживаемый формат файла")
    if not text.strip():
        raise ValueError("Текст не извлечён: пустой документ или PDF-скан; OCR в демо отсутствует")
    if len(text) > 500_000:
        raise ValueError("Превышен предел 500 000 символов")
    return text.strip()
